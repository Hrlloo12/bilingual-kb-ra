from __future__ import annotations

from pathlib import Path

from docx import Document

from rag.parsers.builder import DocumentBuilder, as_bullet
from rag.schemas import ParsedDocument


def parse_docx(path: Path) -> ParsedDocument:
    builder = DocumentBuilder("docx", path.stem)
    for paragraph in Document(str(path)).paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "").lower()
        if style in ("title", "heading 1"):
            builder.set_title(text)
        elif style.startswith("heading"):
            builder.add_heading(text)
        elif style.startswith("list"):
            builder.add_block(as_bullet(text))
        else:
            builder.add_block(text)
    return builder.build()
