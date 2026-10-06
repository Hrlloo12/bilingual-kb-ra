from __future__ import annotations

import json
import re
import shutil
from collections import defaultdict
from pathlib import Path
from typing import Literal

import yaml
from jinja2 import Environment, StrictUndefined
from pydantic import BaseModel, ConfigDict, Field

from rag.facts import Entity, FactBase, FactFile
from rag.normalize import to_eastern_digits

DocLanguage = Literal["ar", "en", "mixed"]
DocFormat = Literal["pdf", "docx", "html", "txt"]

_STANDALONE_DIGITS = re.compile(r"(?<![A-Za-z0-9\-])\d+")
_BLOCK_SEPARATOR = re.compile(r"\n\s*\n")

_HTML_TEMPLATE = """<!doctype html>
<html lang="{{ lang }}" dir="{{ direction }}">
<head>
<meta charset="utf-8">
<title>{{ document.title }}</title>
<style>
@page { size: A4; margin: 2cm; }
body { font-family: "Noto Naskh Arabic", "Noto Sans Arabic", "Noto Sans", "DejaVu Sans", sans-serif; font-size: 11pt; line-height: 1.7; }
h1 { font-size: 18pt; }
h2 { font-size: 14pt; margin-top: 18pt; }
.page-break { break-before: page; }
</style>
</head>
<body>
<h1>{{ document.title }}</h1>
{% for section in document.sections %}<section{% if section.new_page %} class="page-break"{% endif %}>
<h2>{{ section.heading }}</h2>
{% for block in section.blocks %}{% if block.kind == "bullets" %}<ul>
{% for item in block.items %}<li>{{ item }}</li>
{% endfor %}</ul>
{% else %}<p>{{ block.items[0] }}</p>
{% endif %}{% endfor %}</section>
{% endfor %}</body>
</html>
"""


class RepeatSpec(BaseModel):
    model_config = ConfigDict(extra="forbid")
    table: str
    groups: list[str] | None = None


class SectionTemplate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    heading: str = Field(min_length=1)
    body: str = Field(min_length=1)
    new_page: bool = False
    repeat: RepeatSpec | None = None


class EntityView:
    def __init__(self, table: FactFile, entity: Entity) -> None:
        self.key = entity.key
        self.subject = entity.subject
        self.group = entity.group
        self.vars = entity.vars
        self._table = table

    def fact(self, attribute: str) -> str:
        return self._table.entity_fact_id(self.key, attribute)


class DocumentTemplate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    doc_id: str = Field(pattern=r"^[a-z0-9_]+$")
    title: str = Field(min_length=1)
    language: DocLanguage
    format: DocFormat
    domain: str
    digits: Literal["western", "eastern"] = "western"
    sections: list[SectionTemplate] = Field(min_length=1)


class Block(BaseModel):
    kind: Literal["paragraph", "bullets"]
    items: list[str]


class RenderedSection(BaseModel):
    index: int
    heading: str
    blocks: list[Block]
    new_page: bool
    fact_ids: list[str]


class RenderedDocument(BaseModel):
    doc_id: str
    title: str
    language: DocLanguage
    format: DocFormat
    domain: str
    digits: Literal["western", "eastern"]
    source: str
    sections: list[RenderedSection]

    @property
    def is_rtl(self) -> bool:
        return self.language != "en"


def load_templates(directory: Path) -> list[DocumentTemplate]:
    paths = sorted(directory.glob("*.yaml"))
    if not paths:
        raise FileNotFoundError(f"no document templates found in {directory}")
    templates = [DocumentTemplate.model_validate(yaml.safe_load(path.read_text(encoding="utf-8"))) for path in paths]
    seen: set[str] = set()
    for template in templates:
        if template.doc_id in seen:
            raise ValueError(f"duplicate doc_id {template.doc_id}")
        seen.add(template.doc_id)
    return templates


def localize_digits(text: str, digits: str) -> str:
    if digits != "eastern":
        return text
    return _STANDALONE_DIGITS.sub(lambda match: to_eastern_digits(match.group()), text)


def parse_blocks(text: str) -> list[Block]:
    blocks: list[Block] = []
    for raw_block in _BLOCK_SEPARATOR.split(text.strip()):
        lines = [line.strip() for line in raw_block.splitlines() if line.strip()]
        if not lines:
            continue
        if all(line.startswith("- ") for line in lines):
            blocks.append(Block(kind="bullets", items=[line[2:].strip() for line in lines]))
        else:
            blocks.append(Block(kind="paragraph", items=[" ".join(lines)]))
    return blocks


def _section_entities(section: SectionTemplate, fact_base: FactBase) -> list[EntityView | None]:
    if section.repeat is None:
        return [None]
    table = fact_base.entities.get(section.repeat.table)
    if table is None:
        raise KeyError(f"unknown entity table {section.repeat.table}")
    selected = [
        EntityView(table, entity)
        for entity in table.entities
        if section.repeat.groups is None or entity.group in section.repeat.groups
    ]
    if not selected:
        raise ValueError(f"repeat over {section.repeat.table} selected no entities")
    return selected


def render_document(template: DocumentTemplate, fact_base: FactBase) -> RenderedDocument:
    environment = Environment(undefined=StrictUndefined, autoescape=False)
    fact_language = "en" if template.language == "en" else "ar"
    sections: list[RenderedSection] = []
    for section in template.sections:
        for entity in _section_entities(section, fact_base):
            used: set[str] = set()

            def statement(fact_id: str) -> str:
                used.add(fact_id)
                return fact_base.get(fact_id).statement.get(fact_language)

            def value(fact_id: str) -> str:
                used.add(fact_id)
                return fact_base.get(fact_id).value.get(fact_language)

            context = {"s": statement, "v": value, "e": entity}
            heading = environment.from_string(section.heading).render(**context)
            body = environment.from_string(section.body).render(**context)
            sections.append(
                RenderedSection(
                    index=len(sections),
                    heading=localize_digits(heading, template.digits),
                    blocks=parse_blocks(localize_digits(body, template.digits)),
                    new_page=section.new_page,
                    fact_ids=sorted(used),
                )
            )
    headings = [section.heading for section in sections]
    if len(set(headings)) != len(headings):
        raise ValueError(f"{template.doc_id}: section headings must be unique")
    return RenderedDocument(
        doc_id=template.doc_id,
        title=localize_digits(template.title, template.digits),
        language=template.language,
        format=template.format,
        domain=template.domain,
        digits=template.digits,
        source=f"{template.domain}/{template.doc_id}.{template.format}",
        sections=sections,
    )


def check_coverage(documents: list[RenderedDocument], fact_base: FactBase) -> list[str]:
    observed: dict[str, set[str]] = defaultdict(set)
    for document in documents:
        for section in document.sections:
            for fact_id in section.fact_ids:
                observed[fact_id].add(document.language)
    errors = []
    for fact_id, fact in fact_base.facts.items():
        declared = set(fact.coverage)
        actual = observed.get(fact_id, set())
        if actual != declared:
            errors.append(f"{fact_id}: declared coverage {sorted(declared)} but rendered in {sorted(actual)}")
    return errors


def to_html(document: RenderedDocument) -> str:
    environment = Environment(undefined=StrictUndefined, autoescape=True)
    return environment.from_string(_HTML_TEMPLATE).render(
        document=document,
        lang="en" if document.language == "en" else "ar",
        direction="rtl" if document.is_rtl else "ltr",
    )


def to_text(document: RenderedDocument) -> str:
    lines = [f"# {document.title}", ""]
    for section in document.sections:
        lines.extend([f"## {section.heading}", ""])
        for block in section.blocks:
            if block.kind == "bullets":
                lines.extend(f"- {item}" for item in block.items)
            else:
                lines.append(block.items[0])
            lines.append("")
    return "\n".join(lines)


def write_pdf(document: RenderedDocument, path: Path) -> None:
    from weasyprint import HTML

    HTML(string=to_html(document)).write_pdf(str(path))


def write_docx(document: RenderedDocument, path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def apply_direction(paragraph) -> None:
        if not document.is_rtl:
            return
        properties = paragraph._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        properties.append(bidi)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in paragraph.runs:
            run_properties = run._r.get_or_add_rPr()
            rtl = OxmlElement("w:rtl")
            rtl.set(qn("w:val"), "1")
            run_properties.append(rtl)

    docx_document = Document()
    apply_direction(docx_document.add_heading(document.title, level=1))
    for section in document.sections:
        apply_direction(docx_document.add_heading(section.heading, level=2))
        for block in section.blocks:
            if block.kind == "bullets":
                for item in block.items:
                    apply_direction(docx_document.add_paragraph(item, style="List Bullet"))
            else:
                apply_direction(docx_document.add_paragraph(block.items[0]))
    docx_document.save(str(path))


def write_document(document: RenderedDocument, corpus_dir: Path) -> Path:
    path = corpus_dir / document.source
    path.parent.mkdir(parents=True, exist_ok=True)
    if document.format == "pdf":
        write_pdf(document, path)
    elif document.format == "docx":
        write_docx(document, path)
    elif document.format == "html":
        path.write_text(to_html(document), encoding="utf-8")
    else:
        path.write_text(to_text(document), encoding="utf-8")
    return path


def manifest_record(document: RenderedDocument) -> dict:
    return {
        "doc_id": document.doc_id,
        "title": document.title,
        "language": document.language,
        "format": document.format,
        "domain": document.domain,
        "digits": document.digits,
        "source": document.source,
        "sections": [
            {"index": section.index, "heading": section.heading, "fact_ids": section.fact_ids}
            for section in document.sections
        ],
    }


def reset_generated_corpus(corpus_dir: Path) -> None:
    manifest = corpus_dir / "manifest.jsonl"
    if corpus_dir.exists() and any(corpus_dir.iterdir()) and not manifest.exists():
        raise RuntimeError(f"{corpus_dir} contains files that were not generated by render-kb")
    if corpus_dir.exists():
        shutil.rmtree(corpus_dir)
    corpus_dir.mkdir(parents=True)


def write_corpus(documents: list[RenderedDocument], corpus_dir: Path) -> None:
    reset_generated_corpus(corpus_dir)
    for document in documents:
        write_document(document, corpus_dir)
    with (corpus_dir / "manifest.jsonl").open("w", encoding="utf-8") as handle:
        for document in documents:
            handle.write(json.dumps(manifest_record(document), ensure_ascii=False) + "\n")
